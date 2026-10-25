"""Auto-sectioning: split a whole-year student list into sections and lab batches.

The Coordinator uploads one Excel sheet for a year (roll no, name, email, plus any extra columns such
as "Hostel" or "Gender") and chooses the largest section size and the number of lab batches. The
system proposes balanced sections (sizes differ by at most one):

- by roll number: consecutive roll numbers stay together (the usual way), or
- mixed evenly by a column: every section gets a fair share of each value (e.g. hostellers and day
  scholars), each section still sorted by roll number.

Lab batches split each section by roll number. Nothing is saved until the Coordinator confirms.
Saving keeps ids where it can: a section with the same name keeps its id (so its classes stay
linked) and a student with the same roll number or email keeps theirs (so elective lists stay right).
A section that would disappear while it still has classes is refused with a plain message.
"""

from __future__ import annotations

import io
import math
import re

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pydantic import BaseModel

from app.college_edit import dept_slice, merge_dept
from app.io.excel import RowError, slug
from app.models.domain import Batch, College, Section, Student, VenueType

REQUIRED = ["Roll no", "Name", "Email"]
MAX_STUDENTS = 2000


class YearRow(BaseModel):
    row: int
    roll_no: str
    name: str
    email: str
    extra: dict[str, str] = {}


class PlannedSection(BaseModel):
    name: str
    size: int
    batches: list[int]
    first_roll: str
    last_roll: str
    mix: dict[str, int] = {}  # value of the mix column -> students (when mixing)


class SectioningPlan(BaseModel):
    year: int
    shift_id: str
    students: int
    sections: list[PlannedSection]
    warnings: list[str]
    replaces: list[str]  # names of this year's existing sections that the plan replaces


class SectioningOut(BaseModel):
    errors: list[RowError]
    columns: list[str] = []  # extra columns that can be used to mix sections
    plan: SectioningPlan | None = None
    problems: list[str] = []  # reasons it can't be saved
    saved: bool = False


# ---- template + parsing -----------------------------------------------------------------------------

def template() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Students"
    cols = ["Roll no*", "Name*", "Email*", "Hostel", "Gender"]
    ws.append(cols)
    head = PatternFill("solid", fgColor="E0E7FF")
    for i, c in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=i)
        cell.font, cell.fill = Font(bold=True), head
        ws.column_dimensions[cell.column_letter].width = 22
    ws.append(["2026CSE001", "Aarav Sharma", "aarav.sharma@college.edu", "Hosteller", "M"])
    ws.append(["2026CSE002", "Diya Verma", "diya.verma@college.edu", "Day scholar", "F"])
    note = ws.cell(row=1, column=len(cols) + 2, value=(
        "One row per student of the whole year. Roll no, Name and Email are needed (students log in with the email; "
        "first password = roll number). Extra columns like Hostel or Gender are optional: you can ask for every "
        "section to get a fair mix of them. Replace the two example rows."))
    note.alignment, note.font = Alignment(wrap_text=True, vertical="top"), Font(italic=True, color="666666")
    ws.column_dimensions[note.column_letter].width = 60
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def parse(content: bytes) -> tuple[list[YearRow], list[RowError], list[str]]:
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception:  # noqa: BLE001 - any unreadable file gets the same plain answer
        return [], [RowError(sheet="", row=0, message="This file isn't an Excel workbook (.xlsx).")], []
    ws = wb["Students"] if "Students" in wb.sheetnames else wb.worksheets[0]
    sheet = ws.title
    rows_iter = ws.iter_rows(values_only=True)
    header = [str(c or "").replace("*", "").strip() for c in next(rows_iter, [])]
    missing = [c for c in REQUIRED if c.lower() not in {h.lower() for h in header}]
    if missing:
        return [], [RowError(sheet=sheet, row=1, message=f"Missing column(s): {', '.join(missing)}. "
                                                           "Download the year list template to see the format.")], []
    idx = {h.lower(): i for i, h in enumerate(header) if h}
    extras = [h for h in header if h and h.lower() not in {c.lower() for c in REQUIRED}]
    out: list[YearRow] = []
    errors: list[RowError] = []
    seen_roll: dict[str, int] = {}
    seen_email: dict[str, int] = {}
    for r, values in enumerate(rows_iter, start=2):
        cell = {h: (str(values[i]).strip() if i < len(values) and values[i] is not None else "") for h, i in idx.items()}
        if not any(cell.values()):
            continue
        roll, name, email = cell["roll no"], cell["name"], cell["email"].lower()
        if not roll or not name or not email:
            errors.append(RowError(sheet=sheet, row=r, message="Roll no, Name and Email are all needed."))
            continue
        if not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", email):
            errors.append(RowError(sheet=sheet, row=r, message=f"'{email}' isn't a valid email."))
            continue
        if roll.lower() in seen_roll:
            errors.append(RowError(sheet=sheet, row=r, message=f"Roll no {roll} is also on row {seen_roll[roll.lower()]}."))
            continue
        if email in seen_email:
            errors.append(RowError(sheet=sheet, row=r, message=f"Email {email} is also on row {seen_email[email]}."))
            continue
        seen_roll[roll.lower()], seen_email[email] = r, r
        out.append(YearRow(row=r, roll_no=roll, name=name, email=email,
                           extra={h: cell[h.lower()] for h in extras}))
    if not out and not errors:
        errors.append(RowError(sheet=sheet, row=2, message="The sheet has no students."))
    if len(out) > MAX_STUDENTS:
        errors.append(RowError(sheet=sheet, row=0, message=f"At most {MAX_STUDENTS} students per year list."))
    return out, errors, extras


# ---- the split --------------------------------------------------------------------------------------

def _natural(roll: str) -> list:
    """Sort key so that 'CSE2' comes before 'CSE10'."""
    return [int(p) if p.isdigit() else p.lower() for p in re.split(r"(\d+)", roll)]


def _even(n: int, parts: int) -> list[int]:
    base, extra = divmod(n, parts)
    return [base + (1 if i < extra else 0) for i in range(parts)]


def _letters(i: int) -> str:
    out = ""
    i += 1
    while i:
        i, r = divmod(i - 1, 26)
        out = chr(65 + r) + out
    return out


def split(rows: list[YearRow], max_size: int, mix_by: str | None) -> list[list[YearRow]]:
    """Balanced groups of at most ``max_size``, each sorted by roll number."""
    ordered = sorted(rows, key=lambda r: _natural(r.roll_no))
    n = max(1, math.ceil(len(ordered) / max_size))
    if not mix_by:
        groups, start = [], 0
        for size in _even(len(ordered), n):
            groups.append(ordered[start:start + size])
            start += size
        return groups
    # Deal students out like cards, value by value, so each section gets a fair share of every value
    # and the sizes stay within one of each other.
    by_value = sorted(ordered, key=lambda r: (r.extra.get(mix_by, "").lower(), _natural(r.roll_no)))
    groups = [[] for _ in range(n)]
    for i, r in enumerate(by_value):
        groups[i % n].append(r)
    return [sorted(g, key=lambda r: _natural(r.roll_no)) for g in groups]


def plan(college: College, dept: str, year: int, rows: list[YearRow], max_size: int, batches: int,
         shift_id: str | None, mix_by: str | None, prefix: str | None = None) -> tuple[SectioningPlan, list[list[YearRow]]]:
    d = next(x for x in college.departments if x.id == dept)
    prefix = (prefix or f"{d.short}-{year}").strip()
    shift = shift_id or next((s.shift_id for s in college.sections if s.department_id == dept and s.year == year),
                             college.shifts[0].id if college.shifts else "")
    groups = split(rows, max_size, mix_by)
    sections: list[PlannedSection] = []
    for i, g in enumerate(groups):
        mix: dict[str, int] = {}
        if mix_by:
            for r in g:
                v = r.extra.get(mix_by, "") or "(blank)"
                mix[v] = mix.get(v, 0) + 1
        sections.append(PlannedSection(name=f"{prefix}{_letters(i)}", size=len(g),
                                       batches=_even(len(g), batches) if batches > 1 else [],
                                       first_roll=g[0].roll_no if g else "", last_roll=g[-1].roll_no if g else "",
                                       mix=dict(sorted(mix.items()))))
    warnings: list[str] = []
    biggest = max((s.size for s in sections), default=0)
    lecture_rooms = [v for v in college.venues if v.type != VenueType.lab]
    own = [v for v in lecture_rooms if v.department_id == dept]
    if lecture_rooms and biggest > max(v.capacity for v in lecture_rooms):
        warnings.append(f"No lecture room in the college holds {biggest} students. Lower the section size or add a hall.")
    elif own and biggest > max(v.capacity for v in own):
        warnings.append(f"Sections of {biggest} are bigger than every {d.short} room "
                        f"(largest {max(v.capacity for v in own)} seats): their lectures will use shared rooms.")
    labs = [v for v in college.venues if v.type == VenueType.lab]
    biggest_batch = max((b for s in sections for b in s.batches), default=biggest)
    if labs and biggest_batch > max(v.capacity for v in labs):
        warnings.append(f"Lab groups of {biggest_batch} are bigger than every lab (largest {max(v.capacity for v in labs)} "
                        "seats). Use more lab batches.")
    replaces = [s.name for s in college.sections if s.department_id == dept and s.year == year]
    return SectioningPlan(year=year, shift_id=shift, students=len(rows), sections=sections, warnings=warnings,
                          replaces=replaces), groups


def apply(college: College, dept: str, p: SectioningPlan, groups: list[list[YearRow]],
          actor: str) -> tuple[College | None, list[str]]:
    """Replace this year's sections and students with the plan, keeping ids where names / rolls match."""
    data = dept_slice(college, dept)
    old_year = [s for s in data.sections if s.year == p.year]
    old_by_name = {s.name.lower(): s for s in old_year}
    used_batches = {a.batch_id for a in college.allotments if a.batch_id}
    old_students = {st.roll_no.lower(): st for st in college.students}
    old_emails = {st.email.lower(): st for st in college.students}
    problems: list[str] = []
    new_sections: list[Section] = []
    for ps in p.sections:
        prev = old_by_name.get(ps.name.lower())
        sid = prev.id if prev else slug(ps.name)
        batches = [Batch(id=f"{sid}-b{j + 1}", name=f"B{j + 1}", strength=size) for j, size in enumerate(ps.batches)]
        if prev:
            for b in prev.batches:
                if b.id in used_batches and b.id not in {x.id for x in batches}:
                    problems.append(f"{ps.name} has lab classes for batch {b.name}; keep {len(prev.batches)} lab "
                                    "batches or remove those classes first.")
        new_sections.append(Section(id=sid, name=ps.name, department_id=dept, year=p.year, strength=ps.size,
                                    shift_id=p.shift_id, batches=batches))
    kept = {s.id for s in new_sections}
    for s in old_year:
        if s.id not in kept and any(s.id in a.section_ids for a in college.allotments):
            problems.append(f"{s.name} would disappear but still has classes. Keep at least as many sections, or "
                            "remove its classes first.")
    other = {s.id for s in college.sections} - {s.id for s in old_year}
    taken = {s.id for s in new_sections} & other
    if taken:
        problems.append(f"Section name(s) already used elsewhere: {', '.join(sorted(taken))}. Choose another prefix.")
    if problems:
        return None, problems

    students: list[Student] = []
    for sec, g in zip(new_sections, groups):
        starts = [0]
        for b in sec.batches:
            starts.append(starts[-1] + b.strength)
        for k, r in enumerate(g):
            prev = old_students.get(r.roll_no.lower()) or old_emails.get(r.email)
            batch = next((b.id for j, b in enumerate(sec.batches) if starts[j] <= k < starts[j + 1]), None)
            students.append(Student(id=prev.id if prev else f"st-{slug(r.roll_no)}", roll_no=r.roll_no, name=r.name,
                                    email=r.email, section_id=sec.id, batch_id=batch))
    old_year_ids = {s.id for s in old_year}
    data.sections = [s for s in data.sections if s.year != p.year] + new_sections
    data.students = [st for st in data.students if st.section_id not in old_year_ids] + students
    return merge_dept(college, dept, data, actor)

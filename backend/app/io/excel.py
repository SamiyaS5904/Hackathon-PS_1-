"""Excel template (download) and import (upload) for one department's data.

The template is pre-filled with the department's current data, so a Coordinator can
download, edit in Excel and upload again. Import checks every row and reports errors
as "Sheet row N: what's wrong"; nothing is saved until the preview is accepted.
Existing ids are kept by matching names / codes, so the timetable stays linked.
"""

from __future__ import annotations

import io
import re

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from pydantic import BaseModel

from app.college_edit import DeptData, dept_slice
from app.models.domain import (
    Allotment,
    Batch,
    College,
    Kind,
    Section,
    Student,
    Subject,
    Teacher,
    TeacherSubject,
    Venue,
    VenueType,
)

# sheet -> (columns, help line); a * marks a required column
SHEETS: dict[str, tuple[list[str], str]] = {
    "Sections": (["Section*", "Year*", "Students*", "Shift*", "Lab batches"],
                 "Lab batches: e.g. 'B1:30, B2:30' (name:students), or leave empty if labs are taken by the whole section."),
    "Teachers": (["Name*", "Email", "Shifts*", "Max periods per day"],
                 "Shifts: a shift name, or several separated by commas. Max periods per day defaults to 5."),
    "Subjects": (["Code*", "Name*", "Year*", "Lectures per week*", "Labs per week", "Lab equipment"],
                 "Labs per week counts lab blocks (each block = the college's lab length). Lab equipment must match a lab room's equipment."),
    "Can teach": (["Teacher*", "Subject code*", "Preference"],
                  "Which subjects each teacher can teach. Preference 1 = most wanted (default 2). Used for teacher suggestions."),
    "Classes": (["Subject code*", "Sections*", "Type*", "Batch", "Teacher"],
                "One row per class. Several sections separated by commas = one combined class. Type Lab + Batch 'all' = one lab per batch. Teacher may be empty (use suggestions)."),
    "Students": (["Roll no*", "Name*", "Email*", "Section*", "Batch"],
                 "Students log in with this email; their first password is their roll number."),
    "Rooms": (["Name*", "Type*", "Seats*", "Equipment"],
              "Type: Lecture, Lab or Seminar. Equipment for labs, e.g. computers (several separated by commas)."),
}


class RowError(BaseModel):
    sheet: str
    row: int
    message: str


class ImportResult(BaseModel):
    errors: list[RowError]
    counts: dict[str, int]
    data: DeptData | None = None


def slug(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", str(text).lower()).strip("-") or "x"


# ---------------------------------------------------------------------------
# Template
# ---------------------------------------------------------------------------

def _batches_text(sec: Section) -> str:
    return ", ".join(f"{b.name}:{b.strength}" for b in sec.batches)


def template(college: College, dept: str) -> bytes:
    data = dept_slice(college, dept)
    shift_name = {s.id: s.name for s in college.shifts}
    subj = {s.id: s for s in college.subjects}
    teacher = {t.id: t.name for t in college.teachers}
    section = {s.id: s for s in college.sections}

    rows: dict[str, list[list]] = {
        "Sections": [[s.name, s.year, s.strength, shift_name.get(s.shift_id, ""), _batches_text(s)] for s in data.sections],
        "Teachers": [[t.name, "", ", ".join(shift_name.get(x, x) for x in t.shift_ids), t.max_per_day]
                     for t in data.teachers],
        "Subjects": [[s.code, s.name, s.year, s.lectures_per_week, s.labs_per_week or "", s.lab_tag or ""]
                     for s in data.subjects],
        "Can teach": [[teacher.get(ts.teacher_id, ""), subj[ts.subject_id].code if ts.subject_id in subj else "",
                       ts.preference] for ts in data.teacher_subjects],
        "Classes": [],
        "Students": [[st.roll_no, st.name, st.email, section[st.section_id].name,
                      next((b.name for b in section[st.section_id].batches if b.id == st.batch_id), "")]
                     for st in data.students],
        "Rooms": [[v.name, v.type.value.capitalize(), v.capacity, ", ".join(v.tags)] for v in data.venues],
    }
    for a in data.allotments:
        secs = ", ".join(section[i].name for i in a.section_ids)
        batch = next((b.name for b in section[a.section_ids[0]].batches if b.id == a.batch_id), "") if a.batch_id else ""
        rows["Classes"].append([subj[a.subject_id].code if a.subject_id in subj else a.subject_id, secs,
                                "Lab" if a.kind == Kind.lab else "Lecture", batch, teacher.get(a.teacher_id, "")])
    if not data.sections:  # empty department: show the format with example rows
        first_shift = college.shifts[0].name if college.shifts else "Day"
        rows["Sections"] = [["CSE-2A", 2, 60, first_shift, "B1:30, B2:30"]]
        rows["Teachers"] = [["Dr. Harpreet Kaur", "harpreet@college.edu", first_shift, 5]]
        rows["Subjects"] = [["CS201", "Data Structures", 2, 3, 1, "computers"]]
        rows["Can teach"] = [["Dr. Harpreet Kaur", "CS201", 1]]
        rows["Classes"] = [["CS201", "CSE-2A", "Lecture", "", ""], ["CS201", "CSE-2A", "Lab", "all", ""]]
        rows["Students"] = [["CSE2A01", "Aarav Sharma", "cse2a.01@college.edu", "CSE-2A", "B1"]]
        rows["Rooms"] = [["LH-101", "Lecture", 66, ""], ["CSE Lab 1", "Lab", 32, "computers"]]

    wb = Workbook()
    wb.remove(wb.active)
    head = PatternFill("solid", fgColor="E0E7FF")
    for name, (cols, help_text) in SHEETS.items():
        ws = wb.create_sheet(name)
        ws.append(cols)
        for i, _ in enumerate(cols, start=1):
            cell = ws.cell(row=1, column=i)
            cell.font, cell.fill = Font(bold=True), head
            ws.column_dimensions[cell.column_letter].width = max(16, len(cols[i - 1]) + 6)
        for r in rows[name]:
            ws.append(r)
        ws.freeze_panes = "A2"
        note = ws.cell(row=1, column=len(cols) + 2, value=help_text)
        note.alignment, note.font = Alignment(wrap_text=True, vertical="top"), Font(italic=True, color="666666")
        ws.column_dimensions[note.column_letter].width = 60
        if name == "Classes":
            dv = DataValidation(type="list", formula1='"Lecture,Lab"', allow_blank=False)
            ws.add_data_validation(dv)
            dv.add("C2:C1000")
        if name == "Rooms":
            dv = DataValidation(type="list", formula1='"Lecture,Lab,Seminar"', allow_blank=False)
            ws.add_data_validation(dv)
            dv.add("B2:B1000")
        if name in ("Sections", "Teachers") and college.shifts:
            dv = DataValidation(type="list", formula1='"' + ",".join(s.name for s in college.shifts) + '"',
                                allow_blank=False)
            dv.error, dv.showErrorMessage = "Pick a shift (several shifts: type them separated by commas)", False
            ws.add_data_validation(dv)
            dv.add("D2:D1000" if name == "Sections" else "C2:C1000")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Import
# ---------------------------------------------------------------------------

def _rows(wb, sheet: str):
    """(row number, {column: value}) for non-empty rows; columns without the '*' marker."""
    if sheet not in wb.sheetnames:
        return []
    ws = wb[sheet]
    header = [str(c.value or "").replace("*", "").strip() for c in ws[1]]
    out = []
    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        values = {header[j]: (str(v).strip() if v is not None else "") for j, v in enumerate(row) if j < len(header)
                  and header[j]}
        if any(values.values()):
            out.append((i, values))
    return out


def _int(v: str) -> int | None:
    try:
        f = float(v)
    except (TypeError, ValueError):
        return None
    return int(f) if f == int(f) else None


def parse(college: College, dept: str, content: bytes) -> ImportResult:
    errors: list[RowError] = []

    def err(sheet: str, row: int, msg: str) -> None:
        errors.append(RowError(sheet=sheet, row=row, message=msg))

    try:
        wb = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except Exception:  # not an Excel file
        return ImportResult(errors=[RowError(sheet="", row=0, message="This isn't an Excel (.xlsx) file.")], counts={})

    old = dept_slice(college, dept)
    shift_by_name = {s.name.lower(): s.id for s in college.shifts} | {s.id.lower(): s.id for s in college.shifts}
    old_section = {s.name.lower(): s for s in old.sections}
    old_teacher = {t.name.lower(): t for t in old.teachers}
    old_subject = {s.code.lower(): s for s in old.subjects}
    other_teachers = {t.name.lower(): t for t in college.teachers if t.department_id != dept}
    other_subjects = {s.code.lower(): s for s in college.subjects if s.department_id != dept}
    old_venue = {v.name.lower(): v for v in old.venues}
    old_student = {st.email.lower(): st for st in old.students} | {st.roll_no.lower(): st for st in old.students}

    # Rooms
    venues: list[Venue] = []
    for r, v in _rows(wb, "Rooms"):
        name, typ, seats = v.get("Name", ""), v.get("Type", "").lower(), _int(v.get("Seats", ""))
        if not name:
            err("Rooms", r, "Name is missing.")
            continue
        if typ not in ("lecture", "lab", "seminar"):
            err("Rooms", r, f"Type must be Lecture, Lab or Seminar (got '{v.get('Type', '')}').")
            continue
        if not seats or seats < 1:
            err("Rooms", r, "Seats must be a whole number above 0.")
            continue
        vid = old_venue[name.lower()].id if name.lower() in old_venue else f"{dept}-{slug(name)}"
        tags = [slug(t) for t in v.get("Equipment", "").split(",") if t.strip()]
        venues.append(Venue(id=vid, name=name, type=VenueType(typ), capacity=seats, tags=tags, department_id=dept))

    # Sections (+ batches)
    sections: list[Section] = []
    sec_by_name: dict[str, Section] = {}
    for r, v in _rows(wb, "Sections"):
        name, year, strength = v.get("Section", ""), _int(v.get("Year", "")), _int(v.get("Students", ""))
        shift = shift_by_name.get(v.get("Shift", "").lower())
        if not name:
            err("Sections", r, "Section name is missing.")
            continue
        if name.lower() in sec_by_name:
            err("Sections", r, f"Section {name} appears twice.")
            continue
        if not year or year < 1:
            err("Sections", r, "Year must be a whole number like 1, 2, 3.")
            continue
        if not strength or strength < 1:
            err("Sections", r, "Students must be a whole number above 0.")
            continue
        if shift is None:
            err("Sections", r, f"Unknown shift '{v.get('Shift', '')}'. Use one of: "
                f"{', '.join(s.name for s in college.shifts)}.")
            continue
        sid = old_section[name.lower()].id if name.lower() in old_section else slug(name)
        batches: list[Batch] = []
        ok = True
        for part in [p.strip() for p in v.get("Lab batches", "").split(",") if p.strip()]:
            if ":" not in part or _int(part.split(":", 1)[1]) is None:
                err("Sections", r, f"Lab batches must look like 'B1:30, B2:30' (got '{part}').")
                ok = False
                break
            bname, bsize = part.split(":", 1)
            batches.append(Batch(id=f"{sid}-{slug(bname)}", name=bname.strip(), strength=_int(bsize)))
        if not ok:
            continue
        if batches and sum(b.strength for b in batches) != strength:
            err("Sections", r, f"Lab batch sizes add up to {sum(b.strength for b in batches)}, not {strength}.")
            continue
        sec = Section(id=sid, name=name, department_id=dept, year=year, strength=strength, shift_id=shift,
                      batches=batches)
        sections.append(sec)
        sec_by_name[name.lower()] = sec

    # Teachers
    teachers: list[Teacher] = []
    teacher_by_name: dict[str, Teacher] = {}
    for r, v in _rows(wb, "Teachers"):
        name = v.get("Name", "")
        if not name:
            err("Teachers", r, "Name is missing.")
            continue
        if name.lower() in teacher_by_name:
            err("Teachers", r, f"{name} appears twice.")
            continue
        raw = [x.strip() for x in v.get("Shifts", "").split(",") if x.strip()]
        if [x.lower() for x in raw] == ["both"]:
            shift_ids = [s.id for s in college.shifts]
        else:
            shift_ids = [shift_by_name.get(x.lower()) for x in raw]
        if not shift_ids or None in shift_ids:
            err("Teachers", r, f"Shifts must be shift names ({', '.join(s.name for s in college.shifts)}) or 'Both'.")
            continue
        per_day = _int(v.get("Max periods per day", "")) if v.get("Max periods per day") else 5
        if not per_day or per_day < 1:
            err("Teachers", r, "Max periods per day must be a whole number above 0.")
            continue
        prev = old_teacher.get(name.lower())
        t = Teacher(id=prev.id if prev else f"t-{dept}-{slug(name)}", name=name, department_id=dept,
                    shift_ids=shift_ids, max_per_day=per_day, blocked=prev.blocked if prev else [])
        teachers.append(t)
        teacher_by_name[name.lower()] = t

    # Subjects
    subjects: list[Subject] = []
    subj_by_code: dict[str, Subject] = {}
    for r, v in _rows(wb, "Subjects"):
        code, name, year = v.get("Code", ""), v.get("Name", ""), _int(v.get("Year", ""))
        lec = _int(v.get("Lectures per week", "") or "0")
        lab = _int(v.get("Labs per week", "") or "0")
        if not code or not name:
            err("Subjects", r, "Code and Name are required.")
            continue
        if code.lower() in subj_by_code:
            err("Subjects", r, f"Subject code {code} appears twice.")
            continue
        if not year or lec is None or lab is None or lec < 0 or lab < 0:
            err("Subjects", r, "Year, lectures per week and labs per week must be whole numbers.")
            continue
        if lec > len(college.days) or lab > len(college.days):
            err("Subjects", r, f"At most one lecture and one lab a day: {len(college.days)} working days.")
            continue
        tag = slug(v["Lab equipment"]) if v.get("Lab equipment") else None
        prev = old_subject.get(code.lower())
        sub = Subject(id=prev.id if prev else slug(code), code=code, name=name, department_id=dept, year=year,
                      lectures_per_week=lec, labs_per_week=lab, lab_tag=tag)
        subjects.append(sub)
        subj_by_code[code.lower()] = sub

    def find_teacher(name: str) -> Teacher | None:
        return teacher_by_name.get(name.lower()) or other_teachers.get(name.lower())

    def find_subject(code: str) -> Subject | None:
        return subj_by_code.get(code.lower()) or other_subjects.get(code.lower())

    # Can teach
    can: list[TeacherSubject] = []
    for r, v in _rows(wb, "Can teach"):
        t, sub = teacher_by_name.get(v.get("Teacher", "").lower()), find_subject(v.get("Subject code", ""))
        if t is None:
            err("Can teach", r, f"Teacher '{v.get('Teacher', '')}' isn't in the Teachers sheet.")
            continue
        if sub is None:
            err("Can teach", r, f"Subject code '{v.get('Subject code', '')}' isn't in the Subjects sheet.")
            continue
        pref = _int(v.get("Preference", "") or "2")
        if not pref or pref < 1:
            err("Can teach", r, "Preference must be 1, 2, 3 ...")
            continue
        can.append(TeacherSubject(teacher_id=t.id, subject_id=sub.id, preference=pref))

    # Classes
    allotments: list[Allotment] = []
    for r, v in _rows(wb, "Classes"):
        sub = find_subject(v.get("Subject code", ""))
        names = [n.strip() for n in v.get("Sections", "").split(",") if n.strip()]
        secs = [sec_by_name.get(n.lower()) for n in names]
        kind = v.get("Type", "").lower()
        if sub is None:
            err("Classes", r, f"Subject code '{v.get('Subject code', '')}' isn't in the Subjects sheet.")
            continue
        if not names or None in secs:
            missing = [n for n, s in zip(names, secs) if s is None]
            err("Classes", r, f"Unknown section(s): {', '.join(missing) or '(none)'}. Combined classes with other "
                "departments are set up by the College Admin.")
            continue
        if kind not in ("lecture", "lab"):
            err("Classes", r, "Type must be Lecture or Lab.")
            continue
        teacher_id = None
        if v.get("Teacher"):
            t = find_teacher(v["Teacher"])
            if t is None:
                err("Classes", r, f"Teacher '{v['Teacher']}' isn't in the Teachers sheet.")
                continue
            teacher_id = t.id
        if kind == "lecture":
            if not sub.lectures_per_week:
                err("Classes", r, f"{sub.name} has 0 lectures per week in the Subjects sheet.")
                continue
            allotments.append(Allotment(id=f"c-{dept}-{len(allotments) + 1}", subject_id=sub.id, teacher_id=teacher_id,
                                        kind=Kind.lecture, section_ids=[s.id for s in secs]))
            continue
        if not sub.labs_per_week:
            err("Classes", r, f"{sub.name} has 0 labs per week in the Subjects sheet.")
            continue
        if len(secs) != 1:
            err("Classes", r, "A lab belongs to one section (or one of its batches).")
            continue
        sec, batch = secs[0], v.get("Batch", "").strip()
        if batch.lower() in ("all", "all batches") and sec.batches:
            targets = sec.batches
        elif batch:
            b = next((x for x in sec.batches if x.name.lower() == batch.lower()), None)
            if b is None:
                err("Classes", r, f"{sec.name} has no batch '{batch}'.")
                continue
            targets = [b]
        else:
            targets = [None]
        for b in targets:
            allotments.append(Allotment(id=f"c-{dept}-{len(allotments) + 1}", subject_id=sub.id, teacher_id=teacher_id,
                                        kind=Kind.lab, section_ids=[sec.id], batch_id=b.id if b else None))

    # Students
    students: list[Student] = []
    emails: set[str] = set()
    for r, v in _rows(wb, "Students"):
        roll, name, email = v.get("Roll no", ""), v.get("Name", ""), v.get("Email", "").lower()
        sec = sec_by_name.get(v.get("Section", "").lower())
        if not roll or not name or not email:
            err("Students", r, "Roll no, Name and Email are required.")
            continue
        if not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", email):
            err("Students", r, f"'{email}' isn't a valid email.")
            continue
        if email in emails:
            err("Students", r, f"{email} appears twice.")
            continue
        if sec is None:
            err("Students", r, f"Section '{v.get('Section', '')}' isn't in the Sections sheet.")
            continue
        batch_id = None
        if v.get("Batch"):
            b = next((x for x in sec.batches if x.name.lower() == v["Batch"].lower()), None)
            if b is None:
                err("Students", r, f"{sec.name} has no batch '{v['Batch']}'.")
                continue
            batch_id = b.id
        emails.add(email)
        prev = old_student.get(email) or old_student.get(roll.lower())  # keep ids: elective groups refer to them
        students.append(Student(id=prev.id if prev else f"st-{slug(roll)}", roll_no=roll, name=name, email=email,
                                section_id=sec.id, batch_id=batch_id))

    counts = {"sections": len(sections), "teachers": len(teachers), "subjects": len(subjects), "rooms": len(venues),
              "classes": len(allotments), "can_teach": len(can), "students": len(students)}
    if errors:
        return ImportResult(errors=errors, counts=counts)
    data = DeptData(sections=sections, teachers=teachers, subjects=subjects, venues=venues, allotments=allotments,
                    teacher_subjects=can, students=students)
    return ImportResult(errors=[], counts=counts, data=data)

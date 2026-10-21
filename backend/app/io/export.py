"""Excel and PDF timetables for notice boards: one sheet / page per section, teacher or room.

Both formats are drawn from the same ``Grid`` (rows = periods with breaks, columns = working
days), so a printout always matches what the app shows. Classes that overlap on a day (lab
batches in parallel, elective groups) share one merged cell that lists each of them.
"""

from __future__ import annotations

import hashlib
import io
import re
import time
from dataclasses import dataclass, field

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.models.domain import College, Session, Timetable

# Soft fills that print well in black and white too.
PALETTE = ["DCEBFF", "DDF5E3", "FFF1CC", "FBE0E6", "E9E1FF", "D9F3F3", "FFE5D1", "E6EDD9"]


@dataclass
class Row:
    label: str  # "09:00-09:50" or "Break"
    start: str = ""
    end: str = ""
    is_break: bool = False


@dataclass
class Cell:
    lines: list[str]
    span: int = 1  # rows covered
    fill: str = "FFFFFF"
    lunch: bool = False


@dataclass
class Grid:
    title: str
    subtitle: str
    days: list[str]
    rows: list[Row]
    cells: dict[tuple[int, int], Cell] = field(default_factory=dict)  # (row, day index) -> cell
    covered: set[tuple[int, int]] = field(default_factory=set)  # cells inside a span, not the first one


def _fill(subject_id: str) -> str:
    return PALETTE[int(hashlib.md5(subject_id.encode()).hexdigest(), 16) % len(PALETTE)]


def entity_sessions(sessions: list[Session], view: str, entity_id: str) -> list[Session]:
    if view == "section":
        return [s for s in sessions if entity_id in s.section_ids]
    if view == "teacher":
        return [s for s in sessions if s.teacher_id == entity_id]
    return [s for s in sessions if s.venue_id == entity_id]


def entity_name(college: College, view: str, entity_id: str) -> str:
    if view == "section":
        return college.section(entity_id).name
    if view == "teacher":
        return college.teacher(entity_id).name
    return college.venue(entity_id).name


def _rows(college: College, shift_ids: list[str]) -> list[Row]:
    """Periods of the shifts involved, in time order, with break rows between gaps (single shift only)."""
    periods = sorted({(p.start, p.end) for sid in shift_ids for p in college.shift(sid).periods})
    rows: list[Row] = []
    for i, (start, end) in enumerate(periods):
        if i and len(shift_ids) == 1 and periods[i - 1][1] < start:
            rows.append(Row(label=f"Break {periods[i - 1][1]}-{start}", is_break=True))
        rows.append(Row(label=f"{start}-{end}", start=start, end=end))
    return rows


def _line(college: College, s: Session, view: str) -> str:
    subj = college.subject(s.subject_id)
    head = f"{subj.code} {subj.name}"
    if s.kind.value == "lab":
        head += " (lab)"
    if s.elective_slot_id:
        group = college.elective_group(s.elective_slot_id, s.group_id)
        if group.name and group.name.lower() != subj.name.lower():
            head += f" - {group.name}"
    sections = " + ".join(college.section(i).name for i in s.section_ids)
    if s.batch_id:
        batch = next((b.name for i in s.section_ids for b in college.section(i).batches if b.id == s.batch_id), "")
        if view == "section":
            head += f" [{batch}]"
        else:
            sections += f" {batch}"
    teacher = college.teacher(s.teacher_id).name if s.teacher_id else "No teacher yet"
    venue = college.venue(s.venue_id).name
    rest = {"section": [teacher, venue], "teacher": [sections, venue], "venue": [sections, teacher]}[view]
    return " · ".join([head, *rest])


def build_grid(college: College, tt: Timetable, sessions: list[Session], view: str, entity_id: str) -> Grid:
    mine = entity_sessions(sessions, view, entity_id)
    if view == "section":
        sec = college.section(entity_id)
        shift_ids = [sec.shift_id]
        subtitle = f"{sec.strength} students · {college.shift(sec.shift_id).name} shift"
    else:
        shift_ids = sorted({s.shift_id for s in mine}) or [college.shifts[0].id]
        if view == "teacher":
            periods = sum(s.length for s in mine)
            subtitle = f"{periods} periods a week"
        else:
            v = college.venue(entity_id)
            kind = {"lecture": "Lecture room", "lab": "Lab", "seminar": "Seminar hall"}.get(v.type.value, "Room")
            subtitle = f"{kind} · {v.capacity} seats"
    rows = _rows(college, shift_ids)
    grid = Grid(title=entity_name(college, view, entity_id), subtitle=subtitle, days=list(college.days), rows=rows)
    index = {(r.start, r.end): i for i, r in enumerate(rows) if not r.is_break}

    for d, day in enumerate(grid.days):
        todays = sorted((s for s in mine if s.day == day), key=lambda s: s.start)
        # Group overlapping classes (parallel batches, elective groups) into one merged cell.
        clusters: list[list[Session]] = []
        for s in todays:
            if clusters and s.start < max(x.end for x in clusters[-1]):
                clusters[-1].append(s)
            else:
                clusters.append([s])
        for group in clusters:
            start, end = min(s.start for s in group), max(s.end for s in group)
            covered = [i for (a, b), i in index.items() if a >= start and b <= end]
            if not covered:
                continue
            first, last = min(covered), max(covered)
            lines = []
            for s in group:
                text = _line(college, s, view)
                if (s.start, s.end) != (start, end):
                    text += f" ({s.start}-{s.end})"
                lines.append(text)
            grid.cells[(first, d)] = Cell(lines=lines, span=last - first + 1, fill=_fill(group[0].subject_id))
            grid.covered.update((r, d) for r in range(first + 1, last + 1))

    if view == "section":  # the year's lunch period, chosen by the solver, every day
        sec = college.section(entity_id)
        for lunch in tt.lunches:
            if lunch.shift_id == sec.shift_id and lunch.year == sec.year and (lunch.start, lunch.end) in index:
                r = index[(lunch.start, lunch.end)]
                for d in range(len(grid.days)):
                    if (r, d) not in grid.cells and (r, d) not in grid.covered:
                        grid.cells[(r, d)] = Cell(lines=["Lunch"], fill="F2F2F2", lunch=True)
    return grid


def _status(tt: Timetable) -> str:
    if tt.published_at:
        return "Published " + time.strftime("%d %b %Y", time.localtime(tt.published_at))
    return "Draft (not published yet)"


def _sheet_title(name: str, used: set[str]) -> str:
    base = re.sub(r"[\[\]:*?/\\]", "-", name)[:28] or "Sheet"
    title, n = base, 2
    while title.lower() in used:
        title, n = f"{base[:25]} ({n})", n + 1
    used.add(title.lower())
    return title


def to_xlsx(college: College, tt: Timetable, grids: list[Grid]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    used: set[str] = set()
    for g in grids:
        ws = wb.create_sheet(_sheet_title(g.title, used))
        ws["A1"] = f"{college.name} · {g.title}"
        ws["A1"].font = Font(bold=True, size=14)
        ws["A2"] = f"{g.subtitle} · weekly timetable · {_status(tt)}"
        ws["A2"].font = Font(color="666666")
        top = 4
        ws.cell(top, 1, "Time").font = Font(bold=True)
        for d, day in enumerate(g.days):
            c = ws.cell(top, d + 2, day)
            c.font = Font(bold=True)
            c.alignment = Alignment(horizontal="center")
        for r, row in enumerate(g.rows):
            excel_row = top + 1 + r
            label = ws.cell(excel_row, 1, row.label)
            label.font = Font(bold=row.is_break, color="666666" if row.is_break else "000000")
            label.alignment = Alignment(vertical="top")
            ws.row_dimensions[excel_row].height = 14 if row.is_break else 48
            if row.is_break:
                ws.merge_cells(start_row=excel_row, start_column=1, end_row=excel_row, end_column=len(g.days) + 1)
                label.fill = PatternFill("solid", fgColor="F2F2F2")
                continue
            for d in range(len(g.days)):
                if (r, d) in g.covered:
                    continue
                c = ws.cell(excel_row, d + 2)
                c.border = border
                cell = g.cells.get((r, d))
                if cell is None:
                    continue
                c.value = "\n".join(cell.lines)
                c.fill = PatternFill("solid", fgColor=cell.fill)
                c.font = Font(italic=cell.lunch, color="666666" if cell.lunch else "000000", size=9)
                c.alignment = Alignment(wrap_text=True, vertical="top")
                if cell.span > 1:
                    ws.merge_cells(start_row=excel_row, start_column=d + 2,
                                   end_row=excel_row + cell.span - 1, end_column=d + 2)
        ws.column_dimensions["A"].width = 14
        for d in range(len(g.days)):
            ws.column_dimensions[get_column_letter(d + 2)].width = 30
        ws.freeze_panes = ws.cell(top + 1, 2)
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _hex(color: str) -> colors.Color:
    return colors.HexColor("#" + color)


def to_pdf(college: College, tt: Timetable, grids: list[Grid]) -> bytes:
    out = io.BytesIO()
    page = landscape(A4)
    doc = SimpleDocTemplate(out, pagesize=page, leftMargin=10 * mm, rightMargin=10 * mm, topMargin=10 * mm,
                            bottomMargin=10 * mm, title=f"{college.name} timetable", author="Timetabl.ai")
    title = ParagraphStyle("t", fontName="Helvetica-Bold", fontSize=14, leading=17)
    sub = ParagraphStyle("s", fontName="Helvetica", fontSize=9, leading=11, textColor=colors.HexColor("#666666"))
    cell_style = ParagraphStyle("c", fontName="Helvetica", fontSize=7, leading=8.5)
    head_style = ParagraphStyle("h", fontName="Helvetica-Bold", fontSize=8, leading=10)

    def esc(t: str) -> str:
        return t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    story = []
    width = page[0] - 20 * mm
    for n, g in enumerate(grids):
        if n:
            story.append(PageBreak())
        story += [Paragraph(esc(f"{college.name} · {g.title}"), title),
                  Paragraph(esc(f"{g.subtitle} · weekly timetable · {_status(tt)}"), sub), Spacer(1, 4 * mm)]
        data = [[Paragraph("Time", head_style)] + [Paragraph(day, head_style) for day in g.days]]
        style = [("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#BFBFBF")),
                 ("VALIGN", (0, 0), (-1, -1), "TOP"),
                 ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F2F2F2"))]
        for r, row in enumerate(g.rows):
            y = r + 1
            if row.is_break:
                data.append([Paragraph(esc(row.label), sub)] + [""] * len(g.days))
                style += [("SPAN", (0, y), (-1, y)), ("BACKGROUND", (0, y), (-1, y), colors.HexColor("#F7F7F7"))]
                continue
            line = [Paragraph(esc(row.label), sub)]
            for d in range(len(g.days)):
                cell = g.cells.get((r, d))
                if cell is None or (r, d) in g.covered:
                    line.append("")
                    continue
                line.append(Paragraph("<br/>".join(esc(x) for x in cell.lines), cell_style))
                style.append(("BACKGROUND", (d + 1, y), (d + 1, y + cell.span - 1), _hex(cell.fill)))
                if cell.span > 1:
                    style.append(("SPAN", (d + 1, y), (d + 1, y + cell.span - 1)))
            data.append(line)
        day_w = (width - 24 * mm) / max(1, len(g.days))
        table = Table(data, colWidths=[24 * mm] + [day_w] * len(g.days), repeatRows=1)
        table.setStyle(TableStyle(style))
        story.append(table)
    if not grids:
        story.append(Paragraph("Nothing to show.", sub))
    doc.build(story)
    return out.getvalue()

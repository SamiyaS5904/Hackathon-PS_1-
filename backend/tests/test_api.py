"""API tests: login, tenant isolation, role permissions, generate, edit, one-day changes.

One generated demo timetable is shared by the whole module to keep the suite fast.
"""

import time

import pytest
from fastapi.testclient import TestClient

from app.main import app

PW = "demo1234"
DEMO_CLASSES = 232 + 6  # regular classes + open elective (3 groups x 2 meetings)
WEEKDAY_DATE = {"Mon": "2026-10-19", "Tue": "2026-10-20", "Wed": "2026-10-21", "Thu": "2026-10-22", "Fri": "2026-10-23"}


def login(client, email, password=PW):
    r = client.post("/api/auth/login", json={"email": email, "password": password})
    assert r.status_code == 200, r.text
    return {"Authorization": f"Bearer {r.json()['token']}"}


@pytest.fixture(scope="module")
def client():
    with TestClient(app) as c:
        admin = login(c, "admin@demo.timetabl.ai")
        job = c.post("/api/orgs/demo/timetable/generate", json={"time_limit": 25}, headers=admin).json()
        for _ in range(120):
            st = c.get(f"/api/orgs/demo/jobs/{job['id']}", headers=admin).json()
            if st["status"] != "running":
                break
            time.sleep(1)
        assert st["status"] == "done", st
        yield c


@pytest.fixture(scope="module")
def admin(client):
    return login(client, "admin@demo.timetabl.ai")


def test_login_errors_and_me(client):
    bad = client.post("/api/auth/login", json={"email": "admin@demo.timetabl.ai", "password": "nope"})
    assert bad.status_code == 401
    assert client.get("/api/orgs/demo/college").status_code == 401
    assert client.get("/api/me", headers={"Authorization": "Bearer junk"}).status_code == 401
    me = client.get("/api/me", headers=login(client, "admin@demo.timetabl.ai")).json()
    assert me["memberships"][0]["role"] == "admin" and me["memberships"][0]["org_id"] == "demo"


def test_admin_reads_everything(client, admin):
    out = client.get("/api/orgs/demo/timetable", headers=admin).json()
    assert out["violations"] == [] and len(out["timetable"]["sessions"]) == DEMO_CLASSES
    assert len(client.get("/api/orgs/demo/college", headers=admin).json()["students"]) == 720


def test_new_college_is_isolated(client, admin):
    r = client.post("/api/auth/signup", json={"name": "Asha Rao", "email": "asha@newcollege.edu",
                                               "password": "longenough", "college_name": "New College"})
    assert r.status_code == 201, r.text
    mine = {"Authorization": f"Bearer {r.json()['token']}"}
    org = r.json()["memberships"][0]["org_id"]
    assert org == "new-college"
    # Each admin only reaches their own college; outsiders get a plain 404.
    assert client.get("/api/orgs/demo/timetable", headers=mine).status_code == 404
    assert client.get(f"/api/orgs/{org}/college", headers=admin).status_code == 404
    pre = client.get(f"/api/orgs/{org}/precheck", headers=mine).json()
    assert pre["problems"] and "No classes" in pre["problems"][0]
    sample = client.post(f"/api/orgs/{org}/load-sample", headers=mine)
    assert sample.status_code == 200 and len(sample.json()["sections"]) == 12
    again = client.post("/api/auth/signup", json={"name": "X", "email": "asha@newcollege.edu",
                                                   "password": "longenough", "college_name": "Y"})
    assert again.status_code == 409


def test_publish_flow(client, admin):
    st = login(client, "cse3a.01@demo.timetabl.ai")
    t = login(client, "anita.sharma@demo.timetabl.ai")
    co = login(client, "cse.coordinator@demo.timetabl.ai")
    # Before publishing, teachers and students see nothing yet (with a plain reason).
    r = client.get("/api/orgs/demo/timetable", headers=st)
    assert r.status_code == 404 and "published" in r.json()["detail"]
    assert client.get("/api/orgs/demo/timetable", headers=t).status_code == 404
    assert client.get("/api/orgs/demo/timetable", headers=admin).json()["timetable"]["published_at"] is None
    assert client.post("/api/orgs/demo/timetable/publish", headers=co).status_code == 403  # Admin only
    out = client.post("/api/orgs/demo/timetable/publish", headers=admin)
    assert out.status_code == 200 and out.json()["timetable"]["published_at"] and out.json()["ever_published"]
    assert client.post("/api/orgs/demo/timetable/publish", headers=admin).status_code == 409  # already
    assert client.get("/api/orgs/demo/timetable", headers=st).status_code == 200
    audit = client.get("/api/orgs/demo/audit", headers=admin).json()
    assert audit[0]["source"] == "publish" and not audit[0]["can_undo"]


def test_student_sees_only_own_section_and_cannot_edit(client):
    st = login(client, "cse3a.01@demo.timetabl.ai")
    sessions = client.get("/api/orgs/demo/timetable", headers=st).json()["timetable"]["sessions"]
    assert sessions and all("cse-3a" in s["section_ids"] for s in sessions)
    assert all(s["batch_id"] in (None, "cse-3a-b1") for s in sessions)  # student 01 is in batch B1
    college = client.get("/api/orgs/demo/college", headers=st).json()
    assert [s["email"] for s in college["students"]] == ["cse3a.01@demo.timetabl.ai"]  # nobody else's details
    s = sessions[0]
    assert client.get(f"/api/orgs/demo/timetable/sessions/{s['id']}/options", headers=st).status_code == 403
    assert client.post("/api/orgs/demo/timetable/generate", json={}, headers=st).status_code == 403
    assert client.get("/api/orgs/demo/audit", headers=st).status_code == 403


def test_teacher_one_day_change_own_lecture_only(client, admin):
    all_sessions = client.get("/api/orgs/demo/timetable", headers=admin).json()["timetable"]["sessions"]
    mine = next(s for s in all_sessions if s["teacher_id"] == "t-anita")  # Dr. Anita Sharma
    other = next(s for s in all_sessions if s["teacher_id"] != "t-anita")
    t = login(client, "anita.sharma@demo.timetabl.ai")
    # Permanent change: needs the Coordinator (teachers may look at options and send a request).
    r = client.post("/api/orgs/demo/proposals", headers=t,
                    json={"session_id": mine["id"], "day": mine["day"], "start_period": mine["start_period"]})
    assert r.status_code == 403 and "Coordinator" in r.json()["detail"]
    # Not their lecture at all.
    r = client.post("/api/orgs/demo/overrides", headers=t,
                    json={"session_id": other["id"], "date": WEEKDAY_DATE[other["day"]], "kind": "cancel"})
    assert r.status_code == 403
    # One-day cancel of their own lecture: instant.
    r = client.post("/api/orgs/demo/overrides", headers=t, json={
        "session_id": mine["id"], "date": WEEKDAY_DATE[mine["day"]], "kind": "cancel", "reason": "Conference"})
    assert r.status_code == 201, r.text
    audit = client.get("/api/orgs/demo/audit", headers=admin).json()
    assert audit[0]["source"] == "teacher" and audit[0]["actor"] == "Dr. Anita Sharma"


def test_coordinator_limited_to_own_department(client, admin):
    all_sessions = client.get("/api/orgs/demo/timetable", headers=admin).json()["timetable"]["sessions"]
    cse = next(s for s in all_sessions if s["section_ids"] == ["cse-2a"])
    ece = next(s for s in all_sessions if s["section_ids"] == ["ece-2a"])
    co = login(client, "cse.coordinator@demo.timetabl.ai")
    assert client.get(f"/api/orgs/demo/timetable/sessions/{cse['id']}/options", headers=co).status_code == 200
    assert client.get(f"/api/orgs/demo/timetable/sessions/{ece['id']}/options", headers=co).status_code == 403


def test_invites(client, admin):
    co = login(client, "cse.coordinator@demo.timetabl.ai")
    r = client.post("/api/orgs/demo/invites", headers=co,
                    json={"email": "x@y.edu", "role": "coordinator", "department_id": "cse"})
    assert r.status_code == 403
    r = client.post("/api/orgs/demo/invites", headers=co,
                    json={"email": "x@y.edu", "role": "teacher", "department_id": "ece"})
    assert r.status_code == 403
    inv = client.post("/api/orgs/demo/invites", headers=co,
                      json={"email": "new.teacher@demo.timetabl.ai", "role": "teacher", "department_id": "cse"})
    assert inv.status_code == 201, inv.text
    token = inv.json()["token"]
    assert client.get(f"/api/invites/{token}").json()["role"] == "teacher"
    assert client.post(f"/api/invites/{token}/accept", json={"name": "New Teacher", "password": "short"}).status_code == 422
    acc = client.post(f"/api/invites/{token}/accept", json={"name": "New Teacher", "password": "longenough1"})
    assert acc.status_code == 200 and acc.json()["memberships"][0]["role"] == "teacher"
    assert client.post(f"/api/invites/{token}/accept", json={"password": "longenough1"}).status_code == 422
    names = [m["name"] for m in client.get("/api/orgs/demo/members", headers=admin).json()]
    assert "New Teacher" in names


def test_admin_edit_apply_undo(client, admin):
    sessions = client.get("/api/orgs/demo/timetable", headers=admin).json()["timetable"]["sessions"]
    # "repair" (amber) means other classes must shift; some such moves have no clash-free answer,
    # so try amber slots until one can be repaired, like a user would.
    proposal = target = s = None
    for s in sessions:
        if s["elective_slot_id"]:
            continue
        opts = client.get(f"/api/orgs/demo/timetable/sessions/{s['id']}/options", headers=admin).json()
        for target in [o for o in opts if o["status"] == "repair"][:3]:
            body = {"session_id": s["id"], "day": target["day"], "start_period": target["start_period"]}
            r = client.post("/api/orgs/demo/proposals", json=body, headers=admin)
            if r.status_code == 201:
                proposal = r.json()
                break
            assert r.status_code == 422, r.text  # refused with a reason, never an error
        if proposal:
            break
    assert proposal, "expected at least one repairable move"
    assert proposal["changes"][0]["session_id"] == s["id"]
    # Someone else can't apply your proposal.
    co = login(client, "cse.coordinator@demo.timetabl.ai")
    assert client.post(f"/api/orgs/demo/proposals/{proposal['id']}/apply", headers=co).status_code == 403
    r = client.post(f"/api/orgs/demo/proposals/{proposal['id']}/apply", headers=admin)
    assert r.status_code == 200, r.text
    applied = r.json()
    assert applied["violations"] == []
    moved = next(x for x in applied["timetable"]["sessions"] if x["id"] == s["id"])
    assert (moved["day"], moved["start_period"]) == (target["day"], target["start_period"])
    audit = client.get("/api/orgs/demo/audit", headers=admin).json()
    assert audit[0]["source"] == "drag" and audit[0]["can_undo"]
    undone = client.post("/api/orgs/demo/audit/undo", headers=admin).json()
    back = next(x for x in undone["timetable"]["sessions"] if x["id"] == s["id"])
    assert (back["day"], back["start_period"]) == (s["day"], s["start_period"])
    assert client.post(f"/api/orgs/demo/proposals/{proposal['id']}/apply", headers=admin).status_code == 409


def test_state_survives_restart(client, admin):
    """Timetable and audit are persisted: a fresh load from the database sees them."""
    from app.store import Store

    fresh = Store()
    fresh.load()
    st = fresh.org("demo")
    assert st.timetable is not None and len(st.timetable.sessions) == DEMO_CLASSES
    assert st.audit and st.audit[0].source == "generate"


def test_assistant_changes_are_labelled_in_the_activity_log(client, admin):
    sessions = client.get("/api/orgs/demo/timetable", headers=admin).json()["timetable"]["sessions"]
    for s in sessions:
        if s["elective_slot_id"]:
            continue
        opts = client.get(f"/api/orgs/demo/timetable/sessions/{s['id']}/options", headers=admin).json()
        free = [o for o in opts if o["status"] == "free"]
        if free:
            break
    body = {"session_id": s["id"], "day": free[0]["day"], "start_period": free[0]["start_period"], "via": "chatbot"}
    proposal = client.post("/api/orgs/demo/proposals", json=body, headers=admin).json()
    client.post(f"/api/orgs/demo/proposals/{proposal['id']}/apply", headers=admin)
    audit = client.get("/api/orgs/demo/audit", headers=admin).json()
    assert audit[0]["source"] == "chatbot"
    client.post("/api/orgs/demo/audit/undo", headers=admin)


def test_regenerated_draft_keeps_the_published_one_for_students(client, admin):
    """A new draft doesn't reach teachers and students, and they can't edit the old one meanwhile."""
    from app.store import store

    st = login(client, "cse3a.01@demo.timetabl.ai")
    t = login(client, "anita.sharma@demo.timetabl.ai")
    published = store.org("demo").timetable
    draft = published.model_copy(update={"published_at": None, "sessions": published.sessions[1:]})
    store.set_timetable("demo", draft, action="Generated timetable", source="generate", actor="Admin")
    mine = client.get("/api/orgs/demo/timetable", headers=st).json()["timetable"]
    assert mine["published_at"] == published.published_at  # still the published one
    s = next(x for x in published.sessions if x.teacher_id == "t-anita")
    r = client.post("/api/orgs/demo/overrides", headers=t,
                    json={"session_id": s.id, "date": WEEKDAY_DATE[s.day], "kind": "cancel"})
    assert r.status_code == 409 and "published" in r.json()["detail"]
    # Undoing the regeneration brings the published timetable back as the live one.
    client.post("/api/orgs/demo/audit/undo", headers=admin)
    assert store.org("demo").published is store.org("demo").timetable


def test_exports_excel_and_pdf_with_role_limits(client, admin):
    import io

    from openpyxl import load_workbook

    r = client.get("/api/orgs/demo/export?format=xlsx&view=section", headers=admin)
    assert r.status_code == 200 and "attachment" in r.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(r.content))
    assert len(wb.sheetnames) == 12 and "CSE-3A" in wb.sheetnames
    ws = wb["CSE-3A"]
    assert ws.merged_cells.ranges  # labs span several periods
    text = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
    assert "Lunch" in text  # each section's lunch, per day (the demo has no fixed breaks)
    pdf = client.get("/api/orgs/demo/export?format=pdf&view=teacher", headers=admin)
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")
    room = client.get("/api/orgs/demo/export?format=pdf&view=venue&id=lh-101", headers=admin)
    assert room.status_code in (200, 404)

    st = login(client, "cse3a.01@demo.timetabl.ai")
    own = client.get("/api/orgs/demo/export?format=xlsx&view=section&id=cse-3a", headers=st)
    assert own.status_code == 200
    assert client.get("/api/orgs/demo/export?format=xlsx&view=section&id=cse-3b", headers=st).status_code == 403
    assert client.get("/api/orgs/demo/export?format=xlsx&view=section", headers=st).status_code == 403
    # A student's own sheet shows only their batch's labs.
    mine = " ".join(str(c.value) for row in load_workbook(io.BytesIO(own.content)).active.iter_rows()
                    for c in row if c.value)
    assert "[B2]" not in mine


def test_notifications_reach_the_right_people(client, admin):
    st = login(client, "cse3a.01@demo.timetabl.ai")
    other = login(client, "ece3a.01@demo.timetabl.ai")
    co = login(client, "cse.coordinator@demo.timetabl.ai")
    me = client.get("/api/orgs/demo/notifications", headers=st).json()
    assert any(n["title"] == "Timetable published" for n in me["items"]) and me["unread"] >= 1
    # The person who acted isn't notified about their own change.
    assert not any(n["title"] == "Timetable published"
                   for n in client.get("/api/orgs/demo/notifications", headers=admin).json()["items"])
    client.post("/api/orgs/demo/notifications/seen", headers=st)
    assert client.get("/api/orgs/demo/notifications", headers=st).json()["unread"] == 0

    sessions = client.get("/api/orgs/demo/timetable", headers=admin).json()["timetable"]["sessions"]
    s = next(x for x in sessions if x["section_ids"] == ["cse-3a"] and not x["batch_id"]
             and not x["elective_slot_id"])
    r = client.post("/api/orgs/demo/overrides", headers=admin, json={
        "session_id": s["id"], "date": WEEKDAY_DATE[s["day"]], "kind": "cancel", "reason": "Exam duty"})
    assert r.status_code == 201, r.text
    mine = client.get("/api/orgs/demo/notifications", headers=st).json()
    assert mine["unread"] == 1 and "Exam duty" in mine["items"][0]["body"]
    assert client.get("/api/orgs/demo/notifications", headers=co).json()["items"][0]["title"] == mine["items"][0]["title"]
    assert not any("Exam duty" in n["body"]
                   for n in client.get("/api/orgs/demo/notifications", headers=other).json()["items"])
    client.post("/api/orgs/demo/audit/undo", headers=admin)
    assert client.get("/api/orgs/demo/notifications", headers=st).json()["items"][0]["title"].startswith("Change undone")


def test_teacher_change_request_needs_the_coordinator(client, admin):
    t = login(client, "anita.sharma@demo.timetabl.ai")
    co = login(client, "cse.coordinator@demo.timetabl.ai")
    st = login(client, "cse3a.01@demo.timetabl.ai")
    sessions = client.get("/api/orgs/demo/timetable", headers=admin).json()["timetable"]["sessions"]
    mine = [s for s in sessions if s["teacher_id"] == "t-anita" and not s["elective_slot_id"]]
    # Teachers may look at the options for their own lectures, then ask (not apply).
    req = target = s = None
    for s in mine:
        opts = client.get(f"/api/orgs/demo/timetable/sessions/{s['id']}/options", headers=t)
        assert opts.status_code == 200, opts.text
        for target in [o for o in opts.json() if o["status"] == "free"][:2]:
            r = client.post("/api/orgs/demo/change-requests", headers=t, json={
                "session_id": s["id"], "day": target["day"], "start_period": target["start_period"],
                "reason": "Clinic on that morning"})
            if r.status_code == 201:
                req = r.json()
                break
        if req:
            break
    assert req and req["status"] == "pending"
    assert client.post("/api/orgs/demo/proposals", headers=t, json={
        "session_id": s["id"], "day": target["day"], "start_period": target["start_period"]}).status_code == 403
    assert client.post(f"/api/orgs/demo/change-requests/{req['id']}/review", headers=t).status_code == 403
    assert client.get("/api/orgs/demo/change-requests", headers=st).status_code == 403
    # The Coordinator hears about it, reviews (fresh repair + validator) and applies = approves.
    sections = {x for x in s["section_ids"]}
    depts = {sid.split("-")[0] for sid in sections}
    if depts == {"cse"}:
        reviewer = co
        assert any("asks to move" in n["title"]
                   for n in client.get("/api/orgs/demo/notifications", headers=co).json()["items"])
    else:
        reviewer = admin
    assert any(r["id"] == req["id"] for r in client.get("/api/orgs/demo/change-requests", headers=reviewer).json())
    proposal = client.post(f"/api/orgs/demo/change-requests/{req['id']}/review", headers=reviewer)
    assert proposal.status_code == 201, proposal.text
    applied = client.post(f"/api/orgs/demo/proposals/{proposal.json()['id']}/apply", headers=reviewer).json()
    assert applied["violations"] == []
    mine_now = client.get("/api/orgs/demo/change-requests", headers=t).json()
    assert next(r for r in mine_now if r["id"] == req["id"])["status"] == "approved"
    assert client.get("/api/orgs/demo/notifications", headers=t).json()["items"][0]["title"] == \
        "Your change request was approved"
    audit = client.get("/api/orgs/demo/audit", headers=admin).json()
    assert audit[0]["source"] == "teacher" and "requested by Dr. Anita Sharma" in audit[0]["action"]
    client.post("/api/orgs/demo/audit/undo", headers=admin)

    # A second request gets declined with a note.
    r = client.post("/api/orgs/demo/change-requests", headers=t, json={
        "session_id": s["id"], "day": target["day"], "start_period": target["start_period"]})
    assert r.status_code == 201, r.text
    out = client.post(f"/api/orgs/demo/change-requests/{r.json()['id']}/decline", headers=reviewer,
                      json={"note": "Lab is booked then"})
    assert out.status_code == 200 and out.json()["status"] == "declined"
    assert client.post(f"/api/orgs/demo/change-requests/{r.json()['id']}/decline", headers=reviewer,
                       json={}).status_code == 409


def test_teacher_sees_only_own_lectures(client):
    t = login(client, "anita.sharma@demo.timetabl.ai")
    mine = client.get("/api/orgs/demo/timetable", headers=t).json()["timetable"]["sessions"]
    assert mine and all(s["teacher_id"] == "t-anita" for s in mine)
    assert client.get("/api/orgs/demo/export?format=pdf&view=teacher&id=t-anita", headers=t).status_code == 200
    assert client.get("/api/orgs/demo/export?format=pdf&view=section&id=cse-1a", headers=t).status_code == 403


def test_coordinator_sees_only_own_department_readiness(client, admin):
    co = login(client, "cse.coordinator@demo.timetabl.ai")
    mine = client.get("/api/orgs/demo/readiness", headers=co).json()
    assert [d["department_id"] for d in mine] == ["cse"]
    assert len(client.get("/api/orgs/demo/readiness", headers=admin).json()) == 4


def test_attendance_marking_and_student_view(client, admin):
    import datetime

    t = login(client, "anita.sharma@demo.timetabl.ai")
    classes = client.get("/api/orgs/demo/attendance/my-classes", headers=t).json()
    assert classes, "Dr. Anita has classes in the last week"
    c = classes[0]
    sheet = client.get(f"/api/orgs/demo/attendance/class?session_id={c['session_id']}&date={c['date']}", headers=t).json()
    assert sheet["roster"] and not sheet["marked"] and sheet["can_edit"]
    absent = [sheet["roster"][0]["student_id"]]
    saved = client.put("/api/orgs/demo/attendance/class", headers=t,
                       json={"session_id": c["session_id"], "date": c["date"], "absent": absent})
    assert saved.status_code == 200, saved.text
    assert saved.json()["marked"] and saved.json()["absent"] == absent

    # Roles: students can't open a class sheet; another teacher's class is refused; future dates aren't open.
    college = client.get("/api/orgs/demo/college", headers=admin).json()
    email = {st["id"]: st["email"] for st in college["students"]}
    away = login(client, email[absent[0]])
    there = login(client, email[sheet["roster"][1]["student_id"]])
    assert client.get(f"/api/orgs/demo/attendance/class?session_id={c['session_id']}&date={c['date']}",
                      headers=away).status_code == 403
    sessions = client.get("/api/orgs/demo/timetable", headers=admin).json()["timetable"]["sessions"]
    other = next(s for s in sessions if s["teacher_id"] != "t-anita")
    days = ["Mon", "Tue", "Wed", "Thu", "Fri"]
    monday = datetime.date.today() - datetime.timedelta(days=datetime.date.today().weekday())
    other_date = (monday + datetime.timedelta(days=days.index(other["day"]))).isoformat()
    assert client.get(f"/api/orgs/demo/attendance/class?session_id={other['id']}&date={other_date}",
                      headers=t).status_code == 403
    mine = next(s for s in sessions if s["teacher_id"] == "t-anita")
    future = (monday + datetime.timedelta(days=7 + days.index(mine["day"]))).isoformat()
    r = client.put("/api/orgs/demo/attendance/class", headers=t,
                   json={"session_id": mine["id"], "date": future, "absent": []})
    assert r.status_code == 409 and "day of the class" in r.json()["detail"]

    # Students see their own numbers against the college's rule.
    me = client.get("/api/orgs/demo/attendance/me", headers=away).json()
    assert me["held"] == 1 and me["attended"] == 0 and me["status"] == "low" and me["min_percent"] == 75
    ok = client.get("/api/orgs/demo/attendance/me", headers=there).json()
    assert ok["percent"] == 100.0 and ok["status"] == "ok"

    # The Admin sees the section with below-rule students first; the rule is a setting.
    sec = next(i for i in sheet["sections"].split(" · ")[0].split(" + "))
    sec_id = next(x["id"] for x in college["sections"] if x["name"] == sec)
    rows = client.get(f"/api/orgs/demo/attendance/sections/{sec_id}", headers=admin).json()["students"]
    if any(r["student_id"] == absent[0] for r in rows):
        assert rows[0]["low"]
    assert client.put("/api/orgs/demo/attendance/rule", headers=admin, json={"min_percent": 80}).status_code == 200
    assert client.get("/api/orgs/demo/attendance/me", headers=there).json()["min_percent"] == 80
    client.put("/api/orgs/demo/attendance/rule", headers=admin, json={"min_percent": 75})


def test_attendance_advice_numbers():
    from app.attendance import advice

    assert advice(9, 10, 75)[0] == "ok" and "miss 2 more" in advice(9, 10, 75)[1]  # 9/12 = 75%
    status, text = advice(5, 10, 75)
    assert status == "low" and "next 10 classes" in text  # (5+10)/(10+10) = 75%
    assert advice(0, 0, 75)[0] == "none"

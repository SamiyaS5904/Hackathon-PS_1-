"""GNDEC demo college (app/seed/gndec.py): builds, passes the pre-check, generates with no clashes, JPAL is
only for some final-year CSE students, every section gets lunch every day, and the Excel files round-trip."""

import io
import time
from collections import Counter

import pytest
from openpyxl import load_workbook

from app.college_edit import structural_errors
from app.io.excel import parse
from app.seed.gndec import blank_template, build_gndec_college, demo_logins, department_files
from app.solver.feasibility import check
from app.solver.generate import generate
from app.solver.quality import score
from app.validator.check import validate


@pytest.fixture(scope="module")
def college():
    return build_gndec_college()


@pytest.fixture(scope="module")
def timetable(college):
    t0 = time.perf_counter()
    tt = generate(college, time_limit=60)
    tt.solve_seconds = round(time.perf_counter() - t0, 1)
    print(f"\nGNDEC: {len(tt.sessions)} classes in {tt.solve_seconds} s, status {tt.status}, "
          f"quality {score(college, tt.sessions, tt.lunches).score}")
    return tt


def test_builds_with_realistic_size_and_passes_the_precheck(college):
    assert 28 <= len(college.sections) <= 40
    assert {d.short for d in college.departments} == {"CSE", "IT", "ME", "CE", "AS"}
    for dept in ("cse", "it", "me", "ce"):
        assert {s.year for s in college.sections if s.department_id == dept} == {1, 2, 3, 4}
    assert all(len(s.batches) == 2 and sum(b.strength for b in s.batches) == s.strength for s in college.sections)
    assert len(college.students) == sum(s.strength for s in college.sections)
    assert all(st.email == f"{st.roll_no}@gndec.demo" for st in college.students)
    assert len({st.roll_no for st in college.students}) == len(college.students)
    assert len({t.name for t in college.teachers}) == len(college.teachers)
    assert structural_errors(college) == [] and check(college) == []
    assert all(college.department_status[d.id].ready for d in college.departments)
    # Leave cover: every regular subject has at least two teachers who can teach it.
    can = Counter(ts.subject_id for ts in college.teacher_subjects)
    regular = {a.subject_id for a in college.allotments}
    assert all(can[s] >= 2 for s in regular)
    # First-year maths is a combined class.
    assert any(len(a.section_ids) > 1 for a in college.allotments if a.subject_id == "am101")


def test_jpal_is_only_for_some_final_year_cse_students(college):
    slot = college.elective_slot("cse-d4-jpal")
    d4 = {s.id for s in college.sections if s.department_id == "cse" and s.year == 4}
    assert set(slot.section_ids) == d4
    jpal = next(g for g in slot.groups if g.id == "jpal")
    d4_students = [st for st in college.students if st.section_id in d4]
    assert 20 <= len(jpal.student_ids) < len(d4_students) / 2
    members = Counter(i for g in slot.groups for i in g.student_ids)
    assert set(members) == {st.id for st in d4_students} and set(members.values()) == {1}  # each in one group


def test_generates_with_no_clashes_and_lunch_for_everyone(college, timetable):
    assert timetable.status in ("optimal", "feasible"), timetable.messages
    assert validate(college, timetable.sessions, timetable.lunches) == []
    lunches = Counter((ln.section_id, ln.day) for ln in timetable.lunches)
    assert all(lunches[(s.id, d)] == 1 for s in college.sections for d in college.days)
    jpal = [s for s in timetable.sessions if s.elective_slot_id == "cse-d4-jpal" and s.group_id == "jpal"]
    assert len(jpal) == 1 and college.venue(jpal[0].venue_id).capacity >= 30
    assert timetable.solve_seconds < 120


def test_excel_files_round_trip_and_the_blank_template_is_empty(college):
    files = department_files(college)
    assert list(files) == ["GNDEC_1_AS.xlsx", "GNDEC_2_CSE.xlsx", "GNDEC_3_IT.xlsx", "GNDEC_4_ME.xlsx",
                           "GNDEC_5_CE.xlsx"]
    for name, content in files.items():
        dept = name.split("_")[2].removesuffix(".xlsx").lower()
        result = parse(college, dept, content)
        assert result.errors == [], (name, result.errors[:3])
        assert result.counts["sections"] == sum(1 for s in college.sections if s.department_id == dept)
    wb = load_workbook(io.BytesIO(blank_template()))
    assert wb.sheetnames[0] == "Read me"
    assert all(ws.max_row == 1 for ws in wb.worksheets[1:])  # headings only


def test_a_new_college_can_be_built_from_the_files(college):
    """Upload order as the Read me says (Applied Sciences first), into a college that has only the Setup part
    (time grid, departments, shared halls): every section, student, teacher and class arrives."""
    from app.college_edit import merge_dept

    shared = [v for v in college.venues if v.department_id is None]
    c = college.model_copy(update={"sections": [], "teachers": [], "subjects": [], "allotments": [], "students": [],
                                   "teacher_subjects": [], "elective_slots": [], "venues": shared})
    for name, content in department_files(college).items():
        dept = name.split("_")[2].removesuffix(".xlsx").lower()
        result = parse(c, dept, content)
        assert result.errors == [], (name, result.errors[:3])
        c, errs = merge_dept(c, dept, result.data, actor="test")
        assert errs == [], (name, errs)
    for key in ("sections", "students", "teachers", "allotments", "venues"):
        assert len(getattr(c, key)) == len(getattr(college, key)), key
    assert check(c) == []


def test_demo_logins_cover_every_role_and_section(college):
    logins = demo_logins(college)
    roles = Counter(x["role"] for x in logins)
    assert roles["admin"] == 1 and roles["coordinator"] == len(college.departments)
    assert roles["student"] == len(college.sections) + 1  # one per section, plus a JPAL student
    assert all(x["email"].endswith("@gndec.demo") for x in logins)
